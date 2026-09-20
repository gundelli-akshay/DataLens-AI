"""
tests/test_step13_chat.py - Focused tests for unified AI Chat for PDF/DOCX.

Tests:
1. RAG system prompt constraints (grounded context, no hallucination/outside knowledge).
2. format_rag_context formatting (PDF pages, DOCX paragraphs, snippet headers).
3. generate_rag_answer with empty question (ValueError).
4. generate_rag_answer with empty chunks (grounded fallback, no API call).
5. generate_rag_answer with missing API key (ValueError).
6. generate_rag_answer with mocked Groq client (retrieval context + question -> answer + sources).
7. POST /documents/chat integration test for PDF with page citations.
8. POST /documents/chat integration test for DOCX with paragraph citations.
9. POST /documents/chat auto-indexing when file exists on disk.
10. POST /documents/chat validation (empty question -> 400, empty filename -> 400).
11. POST /documents/chat error handling (Groq missing key -> 503, Groq error -> 502).
"""

import sys
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

BACKEND_DIR = Path(__file__).resolve().parent.parent
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

import docx
import pymupdf
from fastapi.testclient import TestClient

from app.core.config import settings
from app.main import app
from app.services.llm import (
    RAG_SYSTEM_PROMPT,
    format_rag_context,
    generate_rag_answer,
)
from app.services.rag import (
    index_document_data,
    vector_index,
)


class TestStep13LLMRAGService(unittest.TestCase):
    """Unit tests for RAG Q&A LLM service in app.services.llm."""

    def test_rag_system_prompt_strictly_restricts_hallucinations(self):
        """Prompt strictly forbids inventing or extrapolating outside knowledge."""
        self.assertIn("ONLY using information explicitly stated", RAG_SYSTEM_PROMPT)
        self.assertIn("Do NOT invent", RAG_SYSTEM_PROMPT)
        self.assertIn("outside knowledge", RAG_SYSTEM_PROMPT)
        self.assertIn("Page X or Paragraph Y", RAG_SYSTEM_PROMPT)

    def test_format_rag_context(self):
        """Context formatting correctly labels PDF pages and DOCX paragraphs."""
        chunks = [
            {
                "filename": "report.pdf",
                "page_number": 2,
                "paragraph_number": None,
                "text": "Solar energy systems generate clean power.",
            },
            {
                "filename": "manual.docx",
                "page_number": None,
                "paragraph_number": 5,
                "text": "Drip irrigation optimizes water conservation.",
            },
        ]
        context_str = format_rag_context(chunks)
        self.assertIn("report.pdf (Page 2)", context_str)
        self.assertIn("manual.docx (Paragraph 5)", context_str)
        self.assertIn("Solar energy systems", context_str)
        self.assertIn("Drip irrigation optimizes", context_str)

    def test_generate_rag_answer_empty_question(self):
        """Empty or whitespace-only question raises ValueError."""
        with self.assertRaises(ValueError):
            generate_rag_answer("", chunks=[{"text": "Sample text"}])
        with self.assertRaises(ValueError):
            generate_rag_answer("   ", chunks=[{"text": "Sample text"}])

    def test_generate_rag_answer_empty_chunks(self):
        """When no chunks are retrieved, return grounded fallback without calling LLM."""
        res = generate_rag_answer("What is the revenue?", chunks=[])
        self.assertEqual(res["status"], "success")
        self.assertIn("does not contain sufficient information", res["answer"])
        self.assertEqual(res["sources"], [])

    def test_generate_rag_answer_missing_api_key(self):
        """Missing API key raises ValueError when client is None."""
        with patch.object(settings, "groq_api_key", ""):
            with self.assertRaises(ValueError) as ctx:
                generate_rag_answer(
                    "What is solar energy?",
                    chunks=[{"text": "Solar energy is light from the sun.", "page_number": 1}],
                )
            self.assertIn("Groq API key is not configured", str(ctx.exception))

    def test_generate_rag_answer_with_mock_client(self):
        """Groq receives only retrieved context + question and returns answer with sources."""
        mock_client = MagicMock()
        mock_choice = MagicMock()
        mock_choice.message.content = "Solar photovoltaics produce clean electricity (Page 1)."
        mock_response = MagicMock()
        mock_response.choices = [mock_choice]
        mock_client.chat.completions.create.return_value = mock_response

        chunks = [
            {
                "filename": "energy.pdf",
                "page_number": 1,
                "paragraph_number": None,
                "text": "Solar photovoltaics and wind turbines produce clean electricity.",
                "score": 0.88,
            },
            {
                "filename": "energy.pdf",
                "page_number": 1,
                "paragraph_number": None,
                "text": "Solar cells convert sunlight directly into electric current.",
                "score": 0.82,
            },
        ]

        result = generate_rag_answer(
            question="How do solar photovoltaics work?",
            chunks=chunks,
            client=mock_client,
        )

        self.assertEqual(result["status"], "success")
        self.assertEqual(result["answer"], "Solar photovoltaics produce clean electricity (Page 1).")
        self.assertEqual(len(result["sources"]), 1)  # Deduplicated by (filename, page, para)
        self.assertEqual(result["sources"][0]["source"], "Page 1")
        self.assertEqual(result["sources"][0]["page_number"], 1)
        self.assertEqual(result["sources"][0]["filename"], "energy.pdf")

        # Verify Groq was called with prompt containing retrieved context & question
        call_kwargs = mock_client.chat.completions.create.call_args[1]
        self.assertEqual(call_kwargs["model"], settings.groq_model)
        user_message = call_kwargs["messages"][1]["content"]
        self.assertIn("Solar photovoltaics and wind turbines", user_message)
        self.assertIn("How do solar photovoltaics work?", user_message)


class TestStep13DocumentChatEndpoint(unittest.TestCase):
    """Integration tests for POST /documents/chat endpoint."""

    def setUp(self):
        self.client = TestClient(app)
        self.temp_files: list[Path] = []
        vector_index.clear()

    def tearDown(self):
        vector_index.clear()
        for p in self.temp_files:
            if p.exists():
                try:
                    p.unlink()
                except Exception:
                    pass

    def _create_sample_pdf(self) -> Path:
        pdf_path = settings.upload_dir_path / "chat_sample_doc.pdf"
        doc = pymupdf.open()

        page1 = doc.new_page()
        page1.insert_text(
            (50, 72),
            "Chapter 1: Renewable Energy Systems\nSolar photovoltaics and wind turbines generate clean electricity with minimal carbon footprint.\nGrid-scale lithium batteries provide overnight backup.",
        )

        page2 = doc.new_page()
        page2.insert_text(
            (50, 72),
            "Chapter 2: Deep Sea Marine Biology\nBioluminescent jellyfish inhabit the abyssal zone where sunlight cannot penetrate.\nHydrothermal vents support chemosynthetic bacteria communities.",
        )

        page3 = doc.new_page()
        page3.insert_text(
            (50, 72),
            "Chapter 3: Quantum Computing Architecture\nSuperconducting qubits operate at cryogenic temperatures near absolute zero millikelvin.\nQuantum error correction codes protect fragile superpositions.",
        )

        doc.save(str(pdf_path))
        doc.close()
        self.temp_files.append(pdf_path)
        return pdf_path

    def _create_sample_docx(self) -> Path:
        docx_path = settings.upload_dir_path / "chat_sample_doc.docx"
        doc = docx.Document()

        doc.add_paragraph(
            "Topic A: Space Telescopes. The James Webb Space Telescope observes primordial infrared galaxy formations."
        )
        doc.add_paragraph(
            "Topic B: Agricultural Irrigation. Drip irrigation systems optimize water delivery directly to root zones, saving 40% freshwater."
        )
        doc.add_paragraph(
            "Topic C: Classical Symphonies. Ludwig van Beethoven composed his Ninth Symphony incorporating Ode to Joy."
        )

        doc.save(str(docx_path))
        self.temp_files.append(docx_path)
        return docx_path

    @patch("app.services.llm.Groq")
    def test_pdf_chat_retrieval_and_answer_with_page_source(self, mock_groq_class):
        """End-to-end PDF chat: retrieval -> Groq -> grounded answer with page source."""
        pdf_path = self._create_sample_pdf()

        # 1. Index document
        index_res = self.client.post("/documents/index", json={"saved_filename": pdf_path.name})
        self.assertEqual(index_res.status_code, 200)

        # 2. Mock Groq client
        mock_groq_instance = MagicMock()
        mock_choice = MagicMock()
        mock_choice.message.content = "Bioluminescent jellyfish live in the abyssal zone where sunlight cannot penetrate (Page 2)."
        mock_response = MagicMock()
        mock_response.choices = [mock_choice]
        mock_groq_instance.chat.completions.create.return_value = mock_response
        mock_groq_class.return_value = mock_groq_instance

        # 3. Chat request
        chat_res = self.client.post(
            "/documents/chat",
            json={
                "filename": pdf_path.name,
                "question": "What organisms inhabit the deep sea abyssal zone?",
                "top_k": 2,
            },
        )
        self.assertEqual(chat_res.status_code, 200)
        data = chat_res.json()
        self.assertEqual(data["status"], "success")
        self.assertIn("Bioluminescent jellyfish", data["answer"])
        self.assertGreater(data["chunks_retrieved"], 0)

        # Sources verify page 2 was retrieved
        sources = data["sources"]
        self.assertGreaterEqual(len(sources), 1)
        self.assertEqual(sources[0]["page_number"], 2)
        self.assertEqual(sources[0]["source"], "Page 2")

    @patch("app.services.llm.Groq")
    def test_docx_chat_retrieval_and_answer_with_paragraph_source(self, mock_groq_class):
        """End-to-end DOCX chat: retrieval -> Groq -> grounded answer with paragraph source."""
        docx_path = self._create_sample_docx()

        # 1. Index document
        index_res = self.client.post("/documents/index", json={"saved_filename": docx_path.name})
        self.assertEqual(index_res.status_code, 200)

        # 2. Mock Groq client
        mock_groq_instance = MagicMock()
        mock_choice = MagicMock()
        mock_choice.message.content = "Drip irrigation optimizes water delivery directly to root zones, saving 40% freshwater."
        mock_response = MagicMock()
        mock_response.choices = [mock_choice]
        mock_groq_instance.chat.completions.create.return_value = mock_response
        mock_groq_class.return_value = mock_groq_instance

        # 3. Chat request
        chat_res = self.client.post(
            "/documents/chat",
            json={
                "filename": docx_path.name,
                "question": "How does drip irrigation save freshwater?",
                "top_k": 2,
            },
        )
        self.assertEqual(chat_res.status_code, 200)
        data = chat_res.json()
        self.assertEqual(data["status"], "success")
        self.assertIn("Drip irrigation", data["answer"])

        # Sources verify paragraph 2 was retrieved
        sources = data["sources"]
        self.assertGreaterEqual(len(sources), 1)
        self.assertEqual(sources[0]["paragraph_number"], 2)
        self.assertEqual(sources[0]["source"], "Paragraph 2")

    @patch("app.services.llm.Groq")
    def test_chat_auto_indexing_from_upload_dir(self, mock_groq_class):
        """If file exists on disk but not yet indexed, /documents/chat indexes it automatically."""
        pdf_path = self._create_sample_pdf()
        # vector_index is empty

        mock_groq_instance = MagicMock()
        mock_choice = MagicMock()
        mock_choice.message.content = "Superconducting qubits operate at cryogenic temperatures near absolute zero (Page 3)."
        mock_response = MagicMock()
        mock_response.choices = [mock_choice]
        mock_groq_instance.chat.completions.create.return_value = mock_response
        mock_groq_class.return_value = mock_groq_instance

        chat_res = self.client.post(
            "/documents/chat",
            json={
                "saved_filename": pdf_path.name,
                "question": "What temperature do superconducting qubits operate at?",
            },
        )
        self.assertEqual(chat_res.status_code, 200)
        data = chat_res.json()
        self.assertEqual(data["status"], "success")
        self.assertIn("Superconducting qubits", data["answer"])
        self.assertEqual(data["sources"][0]["page_number"], 3)

    def test_validation_errors(self):
        """Validation tests for empty question and missing filename."""
        # Empty question
        res1 = self.client.post("/documents/chat", json={"question": "", "filename": "test.pdf"})
        self.assertEqual(res1.status_code, 400)
        self.assertIn("cannot be empty", res1.json()["detail"])

        # Missing filename
        res2 = self.client.post("/documents/chat", json={"question": "What is AI?"})
        self.assertEqual(res2.status_code, 400)
        self.assertIn("filename or saved_filename must be provided", res2.json()["detail"])

    def test_error_handling_unconfigured_api_key(self):
        """Unconfigured Groq API key returns 503 Service Unavailable."""
        pdf_path = self._create_sample_pdf()
        self.client.post("/documents/index", json={"saved_filename": pdf_path.name})

        with patch.object(settings, "groq_api_key", ""):
            res = self.client.post(
                "/documents/chat",
                json={"filename": pdf_path.name, "question": "What is solar energy?"},
            )
            self.assertEqual(res.status_code, 503)
            self.assertIn("Groq API key is not configured", res.json()["detail"])

    @patch("app.services.llm.Groq")
    def test_error_handling_groq_provider_failure(self, mock_groq_class):
        """Provider failure returns 502 Bad Gateway."""
        pdf_path = self._create_sample_pdf()
        self.client.post("/documents/index", json={"saved_filename": pdf_path.name})

        mock_instance = MagicMock()
        mock_instance.chat.completions.create.side_effect = RuntimeError("Groq rate limit exceeded")
        mock_groq_class.return_value = mock_instance

        res = self.client.post(
            "/documents/chat",
            json={"filename": pdf_path.name, "question": "What is solar energy?"},
        )
        self.assertEqual(res.status_code, 502)
        self.assertIn("LLM service error", res.json()["detail"])


if __name__ == "__main__":
    unittest.main()
