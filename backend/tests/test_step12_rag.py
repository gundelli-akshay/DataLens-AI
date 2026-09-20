"""
tests/test_step12_rag.py - Focused tests for Step 12: RAG Pipeline for PDF/DOCX.

Covers:
- Text chunking with overlap and word boundary preservation
- Source reference metadata preservation (PDF page_number, DOCX paragraph_number)
- sentence-transformers embedding generation (dimension, normalization)
- NumPy-based in-memory vector index (cosine similarity, ranking, filtering)
- End-to-end PDF indexing and retrieval with generated test PDF
- End-to-end DOCX indexing and retrieval with generated test DOCX
- API endpoint validation and error handling
"""

import io
from pathlib import Path
import sys
import unittest
import numpy as np

# Ensure backend root is in sys.path
BACKEND_DIR = Path(__file__).resolve().parent.parent
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from fastapi.testclient import TestClient
import pymupdf
import docx

from app.core.config import settings
from app.main import app
from app.services.rag import (
    EmbeddingService,
    InMemoryVectorIndex,
    chunk_extracted_document,
    index_document_data,
    retrieve_relevant_chunks,
    split_text_into_chunks,
    vector_index,
)


class TestChunking(unittest.TestCase):
    """Unit tests for text chunking and document chunking logic."""

    def test_split_text_empty_and_short(self):
        self.assertEqual(split_text_into_chunks(""), [])
        self.assertEqual(split_text_into_chunks("   "), [])

        short_text = "This is a short single sentence."
        chunks = split_text_into_chunks(short_text, chunk_size=100, chunk_overlap=20)
        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0], short_text)

    def test_split_text_overlapping_chunks(self):
        text = (
            "Sentence one discusses data science. Sentence two discusses machine learning. "
            "Sentence three discusses artificial intelligence. Sentence four discusses deep learning models. "
            "Sentence five discusses reinforcement learning and robotic optimization."
        )
        chunk_size = 120
        chunk_overlap = 30
        chunks = split_text_into_chunks(text, chunk_size=chunk_size, chunk_overlap=chunk_overlap)

        self.assertGreater(len(chunks), 1)
        # Check consecutive chunks share overlapping content
        for i in range(len(chunks) - 1):
            curr_words = set(chunks[i].split())
            next_words = set(chunks[i + 1].split())
            common = curr_words.intersection(next_words)
            self.assertTrue(len(common) > 0, f"Chunks {i} and {i+1} should overlap")

    def test_chunk_extracted_pdf_preserves_page_numbers(self):
        doc_data = {
            "filename": "sample_report.pdf",
            "file_type": "PDF",
            "pages": [
                {
                    "page_number": 1,
                    "text": "Page one content about company introduction and overview.",
                },
                {
                    "page_number": 2,
                    "text": "Page two content detailing annual financial metrics and expenses.",
                },
                {
                    "page_number": 3,
                    "text": "Page three content highlighting future projections and strategic roadmap.",
                },
            ],
        }

        chunks = chunk_extracted_document(doc_data, chunk_size=200, chunk_overlap=40)
        self.assertEqual(len(chunks), 3)

        self.assertEqual(chunks[0]["page_number"], 1)
        self.assertIsNone(chunks[0]["paragraph_number"])
        self.assertEqual(chunks[0]["file_type"], "PDF")

        self.assertEqual(chunks[1]["page_number"], 2)
        self.assertIsNone(chunks[1]["paragraph_number"])

        self.assertEqual(chunks[2]["page_number"], 3)
        self.assertIsNone(chunks[2]["paragraph_number"])

    def test_chunk_extracted_docx_preserves_paragraph_numbers(self):
        doc_data = {
            "filename": "handbook.docx",
            "file_type": "DOCX",
            "paragraphs": [
                {
                    "paragraph_number": 1,
                    "text": "Paragraph one: Code of conduct and workplace expectations.",
                },
                {
                    "paragraph_number": 2,
                    "text": "Paragraph two: Leave policies and remote work guidelines.",
                },
            ],
        }

        chunks = chunk_extracted_document(doc_data, chunk_size=200, chunk_overlap=40)
        self.assertEqual(len(chunks), 2)

        self.assertIsNone(chunks[0]["page_number"])
        self.assertEqual(chunks[0]["paragraph_number"], 1)
        self.assertEqual(chunks[0]["file_type"], "DOCX")

        self.assertIsNone(chunks[1]["page_number"])
        self.assertEqual(chunks[1]["paragraph_number"], 2)
        self.assertEqual(chunks[1]["file_type"], "DOCX")


class TestInMemoryVectorIndex(unittest.TestCase):
    """Unit tests for the NumPy vector store."""

    def setUp(self):
        self.index = InMemoryVectorIndex()

    def test_add_and_search(self):
        chunks = [
            {"chunk_id": "c1", "text": "Apple fruit orchard", "filename": "doc1.pdf", "page_number": 1},
            {"chunk_id": "c2", "text": "Automotive car engine", "filename": "doc2.pdf", "page_number": 2},
        ]
        # Vectors where query aligns strongly with c2
        emb_doc1 = np.array([1.0, 0.0, 0.0], dtype=np.float32)
        emb_doc2 = np.array([0.0, 1.0, 0.0], dtype=np.float32)
        embeddings = np.vstack([emb_doc1, emb_doc2])

        self.index.add_documents(chunks, embeddings)
        self.assertEqual(self.index.count(), 2)

        query_vec = np.array([0.1, 0.95, 0.0], dtype=np.float32)
        results = self.index.search(query_vec, top_k=2)

        self.assertEqual(len(results), 2)
        self.assertEqual(results[0]["chunk_id"], "c2")
        self.assertGreater(results[0]["score"], results[1]["score"])

    def test_filter_by_filename(self):
        chunks = [
            {"chunk_id": "c1", "text": "First chunk", "filename": "alpha.pdf"},
            {"chunk_id": "c2", "text": "Second chunk", "filename": "beta.docx"},
        ]
        embeddings = np.array([[1.0, 0.0], [0.0, 1.0]], dtype=np.float32)
        self.index.add_documents(chunks, embeddings)

        results = self.index.search(np.array([1.0, 0.0]), top_k=5, filename="beta.docx")
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]["filename"], "beta.docx")

    def test_clear_index(self):
        chunks = [{"chunk_id": "c1", "text": "Chunk", "filename": "doc.pdf"}]
        embeddings = np.array([[1.0, 0.0]], dtype=np.float32)
        self.index.add_documents(chunks, embeddings)
        self.assertEqual(self.index.count(), 1)

        self.index.clear()
        self.assertEqual(self.index.count(), 0)
        self.assertEqual(self.index.search(np.array([1.0, 0.0])), [])


class TestEmbeddingService(unittest.TestCase):
    """Unit tests for sentence-transformers embedding generation."""

    def test_encode_dimensions_and_normalization(self):
        texts = ["Document retrieval test query.", "Second sentence for embedding."]
        embeddings = EmbeddingService.encode(texts)

        self.assertEqual(embeddings.shape, (2, 384))
        self.assertEqual(embeddings.dtype, np.float32)

        # Verify unit norm due to normalize_embeddings=True
        norm_0 = np.linalg.norm(embeddings[0])
        norm_1 = np.linalg.norm(embeddings[1])
        self.assertAlmostEqual(float(norm_0), 1.0, places=4)
        self.assertAlmostEqual(float(norm_1), 1.0, places=4)


class TestRAGPipelineEndToEnd(unittest.TestCase):
    """Integration tests using generated PDF and DOCX documents."""

    def setUp(self):
        from app.core.auth import get_current_user
        from app.db.models import User
        self.test_user = User(id=1, email="test12@example.com", full_name="Test User", auth_provider="email")
        app.dependency_overrides[get_current_user] = lambda: self.test_user
        self.client = TestClient(app)
        self.temp_files: list[Path] = []
        vector_index.clear()

    def tearDown(self):
        from app.core.auth import get_current_user
        app.dependency_overrides.pop(get_current_user, None)
        vector_index.clear()
        for p in self.temp_files:
            if p.exists():
                try:
                    p.unlink()
                except Exception:
                    pass

    def _create_test_pdf(self) -> Path:
        """Create a 3-page test PDF with distinct semantic topics."""
        pdf_path = settings.upload_dir_path / "test_rag_sample.pdf"
        doc = pymupdf.open()

        page1 = doc.new_page()
        page1.insert_text(
            (50, 72),
            "Chapter 1: Renewable Energy Systems\n"
            "Solar photovoltaics and wind turbines generate clean electricity with minimal carbon footprint.\n"
            "Energy storage solutions such as grid-scale lithium batteries provide overnight backup.",
        )

        page2 = doc.new_page()
        page2.insert_text(
            (50, 72),
            "Chapter 2: Deep Sea Marine Biology\n"
            "Bioluminescent jellyfish inhabit the abyssal zone where sunlight cannot penetrate.\n"
            "Hydrothermal vents support chemosynthetic bacteria communities in extreme pressure.",
        )

        page3 = doc.new_page()
        page3.insert_text(
            (50, 72),
            "Chapter 3: Quantum Computing Architecture\n"
            "Superconducting qubits operate at cryogenic temperatures near absolute zero millikelvin.\n"
            "Quantum error correction codes protect fragile superpositions against thermal decoherence.",
        )

        doc.save(str(pdf_path))
        doc.close()
        self.temp_files.append(pdf_path)
        return pdf_path

    def _create_test_docx(self) -> Path:
        """Create a 3-paragraph test DOCX with distinct semantic topics."""
        docx_path = settings.upload_dir_path / "test_rag_sample.docx"
        doc = docx.Document()

        doc.add_paragraph(
            "Topic A: International Space Exploration. "
            "The James Webb Space Telescope observes primordial infrared galaxy formations from the Sun-Earth L2 orbit."
        )
        doc.add_paragraph(
            "Topic B: Agricultural Irrigation Technology. "
            "Drip irrigation systems optimize water delivery directly to crop root zones, conserving 40% freshwater."
        )
        doc.add_paragraph(
            "Topic C: Classical Symphony Composition. "
            "Ludwig van Beethoven composed his Ninth Symphony incorporating Schiller's Ode to Joy choral finale."
        )

        doc.save(str(docx_path))
        self.temp_files.append(docx_path)
        return docx_path

    def test_pdf_indexing_and_retrieval_with_page_references(self):
        pdf_path = self._create_test_pdf()

        # 1. Index the PDF via API
        index_res = self.client.post(
            "/documents/index",
            json={"saved_filename": pdf_path.name, "chunk_size": 400, "chunk_overlap": 50},
        )
        self.assertEqual(index_res.status_code, 200)
        data = index_res.json()
        self.assertEqual(data["status"], "success")
        self.assertEqual(data["file_type"], "PDF")
        self.assertGreaterEqual(data["chunks_indexed"], 3)

        # 2. Retrieve Query 1 -> Should point to Page 2 (Marine Biology)
        query1 = "bioluminescent organisms deep sea abyssal vents"
        ret_res1 = self.client.post(
            "/documents/retrieve",
            json={"query": query1, "top_k": 2},
        )
        self.assertEqual(ret_res1.status_code, 200)
        results1 = ret_res1.json()["results"]
        self.assertGreaterEqual(len(results1), 1)
        top1 = results1[0]
        self.assertEqual(top1["page_number"], 2)
        self.assertIsNone(top1["paragraph_number"])
        self.assertIn("Bioluminescent", top1["text"])
        self.assertGreater(top1["score"], 0.3)

        # 3. Retrieve Query 2 -> Should point to Page 3 (Quantum Computing)
        query2 = "cryogenic superconducting qubits decoherence"
        ret_res2 = self.client.post(
            "/documents/retrieve",
            json={"query": query2, "top_k": 2},
        )
        self.assertEqual(ret_res2.status_code, 200)
        results2 = ret_res2.json()["results"]
        self.assertGreaterEqual(len(results2), 1)
        top2 = results2[0]
        self.assertEqual(top2["page_number"], 3)
        self.assertIsNone(top2["paragraph_number"])
        self.assertIn("Quantum", top2["text"])

    def test_docx_indexing_and_retrieval_with_paragraph_references(self):
        docx_path = self._create_test_docx()

        # 1. Index the DOCX via API
        index_res = self.client.post(
            "/documents/index",
            json={"saved_filename": docx_path.name, "chunk_size": 300, "chunk_overlap": 40},
        )
        self.assertEqual(index_res.status_code, 200)
        data = index_res.json()
        self.assertEqual(data["status"], "success")
        self.assertEqual(data["file_type"], "DOCX")
        self.assertGreaterEqual(data["chunks_indexed"], 3)

        # 2. Retrieve Query -> Should point to Paragraph 2 (Drip irrigation)
        query = "agricultural water conservation crop root zone"
        ret_res = self.client.post(
            "/documents/retrieve",
            json={"query": query, "top_k": 2},
        )
        self.assertEqual(ret_res.status_code, 200)
        results = ret_res.json()["results"]
        self.assertGreaterEqual(len(results), 1)
        top_match = results[0]
        self.assertIsNone(top_match["page_number"])
        self.assertEqual(top_match["paragraph_number"], 2)
        self.assertIn("irrigation", top_match["text"])

        # 3. Retrieve Query -> Should point to Paragraph 3 (Beethoven symphony)
        query_music = "Beethoven Ode to Joy choral symphony"
        ret_res_music = self.client.post(
            "/documents/retrieve",
            json={"query": query_music, "top_k": 2},
        )
        self.assertEqual(ret_res_music.status_code, 200)
        results_music = ret_res_music.json()["results"]
        top_music = results_music[0]
        self.assertEqual(top_music["paragraph_number"], 3)
        self.assertIn("Beethoven", top_music["text"])

    def test_multipart_upload_and_indexing(self):
        pdf_path = self._create_test_pdf()
        content = pdf_path.read_bytes()

        response = self.client.post(
            "/documents/index",
            files={"file": ("direct_upload.pdf", content, "application/pdf")},
            data={"chunk_size": "500", "chunk_overlap": "100"},
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["status"], "success")
        self.assertGreater(data["chunks_indexed"], 0)

    def test_api_validation_and_errors(self):
        # Empty query
        res_empty = self.client.post("/documents/retrieve", json={"query": ""})
        self.assertEqual(res_empty.status_code, 400)

        # Search on empty index
        vector_index.clear()
        res_empty_index = self.client.post("/documents/retrieve", json={"query": "hello world"})
        self.assertEqual(res_empty_index.status_code, 200)
        self.assertEqual(res_empty_index.json()["count"], 0)

        # Index unsupported file type
        res_csv = self.client.post(
            "/documents/index",
            json={"saved_filename": "data.csv"},
        )
        self.assertEqual(res_csv.status_code, 415)

        # Index non-existent file
        res_notfound = self.client.post(
            "/documents/index",
            json={"saved_filename": "non_existent_file.pdf"},
        )
        self.assertEqual(res_notfound.status_code, 404)

        # Clear index endpoint
        res_clear = self.client.post("/documents/clear-index")
        self.assertEqual(res_clear.status_code, 200)
        self.assertEqual(res_clear.json()["total_indexed_chunks"], 0)


if __name__ == "__main__":
    unittest.main()
