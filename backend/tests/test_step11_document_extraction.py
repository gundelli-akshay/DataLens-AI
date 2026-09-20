"""
tests/test_step11_document_extraction.py - Focused tests for Step 11: PDF and DOCX text extraction.
"""

import io
import os
import sys
from pathlib import Path
import unittest

# Ensure backend directory is in path
BACKEND_DIR = Path(__file__).resolve().parent.parent
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from fastapi.testclient import TestClient
import pymupdf
import docx

from app.core.config import settings
from app.main import app
from app.services.document_extraction import (
    extract_document,
    extract_text_from_docx,
    extract_text_from_pdf,
)


class TestStep11DocumentExtractionService(unittest.TestCase):
    """Unit tests for document extraction service functions."""

    def setUp(self):
        self.temp_files: list[Path] = []

    def tearDown(self):
        for f in self.temp_files:
            if f.exists():
                try:
                    f.unlink()
                except Exception:
                    pass

    def _create_test_pdf(self, pages_text: list[str], filename: str = "test.pdf") -> Path:
        """Helper to create a temporary PDF with specified pages."""
        doc = pymupdf.open()
        for text in pages_text:
            page = doc.new_page()
            page.insert_text((50, 72), text)
        path = settings.upload_dir_path / filename
        doc.save(str(path))
        doc.close()
        self.temp_files.append(path)
        return path

    def _create_test_docx(self, paragraphs_text: list[str], filename: str = "test.docx") -> Path:
        """Helper to create a temporary DOCX with specified paragraphs."""
        doc = docx.Document()
        for text in paragraphs_text:
            doc.add_paragraph(text)
        path = settings.upload_dir_path / filename
        doc.save(str(path))
        self.temp_files.append(path)
        return path

    def test_extract_pdf_page_by_page_success(self):
        """Extracts text page-by-page from a multi-page PDF."""
        pdf_path = self._create_test_pdf(
            ["Quarterly Financial Report Page 1", "Detailed Operations Overview Page 2"],
            "financial_report.pdf",
        )

        res = extract_text_from_pdf(pdf_path, "financial_report.pdf")

        self.assertEqual(res["status"], "success")
        self.assertEqual(res["file_type"], "PDF")
        self.assertEqual(res["filename"], "financial_report.pdf")
        self.assertEqual(res["page_count"], 2)
        self.assertIsNone(res["paragraph_count"])
        self.assertEqual(len(res["pages"]), 2)

        self.assertEqual(res["pages"][0]["page_number"], 1)
        self.assertIn("Quarterly Financial Report Page 1", res["pages"][0]["text"])

        self.assertEqual(res["pages"][1]["page_number"], 2)
        self.assertIn("Detailed Operations Overview Page 2", res["pages"][1]["text"])

        self.assertIn("Quarterly Financial Report Page 1", res["text"])
        self.assertIn("Detailed Operations Overview Page 2", res["text"])
        self.assertGreater(res["total_characters"], 0)
        self.assertGreater(res["total_words"], 0)

    def test_extract_docx_paragraph_by_paragraph_success(self):
        """Extracts text paragraph-by-paragraph from a DOCX document."""
        docx_path = self._create_test_docx(
            ["Executive Summary paragraph.", "", "Key Findings and strategic recommendations."],
            "project_overview.docx",
        )

        res = extract_text_from_docx(docx_path, "project_overview.docx")

        self.assertEqual(res["status"], "success")
        self.assertEqual(res["file_type"], "DOCX")
        self.assertEqual(res["filename"], "project_overview.docx")
        self.assertIsNone(res["page_count"])
        self.assertEqual(res["paragraph_count"], 2)  # Blank paragraph ignored
        self.assertEqual(len(res["paragraphs"]), 2)

        self.assertEqual(res["paragraphs"][0]["paragraph_number"], 1)
        self.assertEqual(res["paragraphs"][0]["text"], "Executive Summary paragraph.")

        self.assertEqual(res["paragraphs"][1]["paragraph_number"], 2)
        self.assertEqual(res["paragraphs"][1]["text"], "Key Findings and strategic recommendations.")

        self.assertIn("Executive Summary paragraph.", res["text"])
        self.assertIn("Key Findings and strategic recommendations.", res["text"])
        self.assertGreater(res["total_characters"], 0)
        self.assertGreater(res["total_words"], 0)

    def test_extract_document_dispatcher(self):
        """extract_document dispatches correctly and rejects unsupported extensions."""
        pdf_path = self._create_test_pdf(["PDF dispatcher test"], "dispatch.pdf")
        docx_path = self._create_test_docx(["DOCX dispatcher test"], "dispatch.docx")

        pdf_res = extract_document(pdf_path, "dispatch.pdf")
        self.assertEqual(pdf_res["file_type"], "PDF")

        docx_res = extract_document(docx_path, "dispatch.docx")
        self.assertEqual(docx_res["file_type"], "DOCX")

        with self.assertRaises(ValueError) as ctx:
            extract_document(Path("dummy.csv"))
        self.assertIn("Unsupported document format", str(ctx.exception))

    def test_extract_nonexistent_file_raises(self):
        """Raises FileNotFoundError when the requested document does not exist."""
        with self.assertRaises(FileNotFoundError):
            extract_text_from_pdf(settings.upload_dir_path / "nonexistent.pdf")

        with self.assertRaises(FileNotFoundError):
            extract_text_from_docx(settings.upload_dir_path / "nonexistent.docx")

    def test_extract_corrupted_file_raises_value_error(self):
        """Raises ValueError when document file is corrupted."""
        corrupt_pdf = settings.upload_dir_path / "corrupt.pdf"
        corrupt_pdf.write_bytes(b"not a valid pdf content")
        self.temp_files.append(corrupt_pdf)

        with self.assertRaises(ValueError) as ctx:
            extract_text_from_pdf(corrupt_pdf)
        self.assertIn("Unable to read PDF file", str(ctx.exception))


class TestStep11DocumentExtractionEndpoint(unittest.TestCase):
    """Integration tests for POST /documents/extract and /documents/ endpoints."""

    @classmethod
    def setUpClass(cls):
        from app.core.auth import get_current_user
        from app.db.models import User
        cls.test_user = User(id=1, email="test11@example.com", full_name="Test User", auth_provider="email")
        app.dependency_overrides[get_current_user] = lambda: cls.test_user
        cls.client = TestClient(app)

    @classmethod
    def tearDownClass(cls):
        from app.core.auth import get_current_user
        app.dependency_overrides.pop(get_current_user, None)

    def setUp(self):
        self.temp_files: list[Path] = []

    def tearDown(self):
        for f in self.temp_files:
            if f.exists():
                try:
                    f.unlink()
                except Exception:
                    pass

    def _create_test_pdf(self, text: str, filename: str) -> str:
        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_text((50, 72), text)
        path = settings.upload_dir_path / filename
        doc.save(str(path))
        doc.close()
        self.temp_files.append(path)
        return filename

    def _create_test_docx(self, text: str, filename: str) -> str:
        doc = docx.Document()
        doc.add_paragraph(text)
        path = settings.upload_dir_path / filename
        doc.save(str(path))
        self.temp_files.append(path)
        return filename

    def test_endpoint_extract_pdf_success(self):
        """POST /documents/extract processes an uploaded PDF and returns clean text + metadata."""
        saved_name = "00000000000000000000000000000010_annual_report.pdf"
        self._create_test_pdf("Annual revenue grew by 25% year over year.", saved_name)

        resp = self.client.post("/documents/extract", json={"saved_filename": saved_name})
        self.assertEqual(resp.status_code, 200)

        data = resp.json()
        self.assertEqual(data["status"], "success")
        self.assertEqual(data["file_type"], "PDF")
        self.assertEqual(data["filename"], "annual_report.pdf")
        self.assertEqual(data["page_count"], 1)
        self.assertIsNone(data["paragraph_count"])
        self.assertIn("Annual revenue grew by 25%", data["text"])
        self.assertEqual(len(data["pages"]), 1)
        self.assertIn("Annual revenue grew by 25%", data["pages"][0]["text"])

    def test_endpoint_extract_docx_success(self):
        """POST /documents/ processes an uploaded DOCX and returns clean text + metadata."""
        saved_name = "00000000000000000000000000000020_project_brief.docx"
        self._create_test_docx("Project Brief: Next generation AI analytics.", saved_name)

        resp = self.client.post("/documents/", json={"saved_filename": saved_name})
        self.assertEqual(resp.status_code, 200)

        data = resp.json()
        self.assertEqual(data["status"], "success")
        self.assertEqual(data["file_type"], "DOCX")
        self.assertEqual(data["filename"], "project_brief.docx")
        self.assertIsNone(data["page_count"])
        self.assertEqual(data["paragraph_count"], 1)
        self.assertIn("Project Brief: Next generation AI analytics.", data["text"])
        self.assertEqual(len(data["paragraphs"]), 1)
        self.assertIn("Project Brief: Next generation AI analytics.", data["paragraphs"][0]["text"])

    def test_endpoint_direct_multipart_upload_success(self):
        """POST /documents/extract accepts direct multipart file upload."""
        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_text((50, 72), "Uploaded directly via multipart form.")
        pdf_bytes = doc.tobytes()
        doc.close()

        resp = self.client.post(
            "/documents/extract",
            files={"file": ("direct_upload.pdf", pdf_bytes, "application/pdf")},
        )
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertEqual(data["file_type"], "PDF")
        self.assertIn("Uploaded directly via multipart form.", data["text"])

    def test_endpoint_rejects_tabular_csv_cleanly(self):
        """Rejects CSV file cleanly with 415 indicating it is a tabular file."""
        saved_csv = "00000000000000000000000000000030_data.csv"
        csv_path = settings.upload_dir_path / saved_csv
        csv_path.write_text("a,b\n1,2\n", encoding="utf-8")
        self.temp_files.append(csv_path)

        resp = self.client.post("/documents/extract", json={"saved_filename": saved_csv})
        self.assertEqual(resp.status_code, 415)
        self.assertIn("tabular data format", resp.json()["detail"])

    def test_endpoint_rejects_unsupported_extension_cleanly(self):
        """Rejects unsupported file type (e.g. .txt) with 415."""
        saved_txt = "00000000000000000000000000000040_notes.txt"
        txt_path = settings.upload_dir_path / saved_txt
        txt_path.write_text("plain text", encoding="utf-8")
        self.temp_files.append(txt_path)

        resp = self.client.post("/documents/extract", json={"saved_filename": saved_txt})
        self.assertEqual(resp.status_code, 415)
        self.assertIn("not a supported document format", resp.json()["detail"])

    def test_endpoint_missing_parameters(self):
        """Returns 400 when neither saved_filename nor file is provided."""
        resp = self.client.post("/documents/extract", json={})
        self.assertEqual(resp.status_code, 400)
        self.assertIn("Either 'saved_filename' or a document file", resp.json()["detail"])

    def test_endpoint_path_traversal_rejected(self):
        """Returns 400 when directory traversal is attempted."""
        resp = self.client.post("/documents/extract", json={"saved_filename": "../secret.pdf"})
        self.assertEqual(resp.status_code, 400)
        self.assertIn("Invalid filename", resp.json()["detail"])

    def test_endpoint_nonexistent_file_returns_404(self):
        """Returns 404 when file does not exist in uploads."""
        resp = self.client.post("/documents/extract", json={"saved_filename": "non_existent.pdf"})
        self.assertEqual(resp.status_code, 404)
        self.assertIn("was not found", resp.json()["detail"])

    def test_endpoint_corrupted_document_returns_422(self):
        """Returns 422 when document file is invalid/corrupted."""
        corrupt_name = "00000000000000000000000000000050_corrupt.pdf"
        corrupt_path = settings.upload_dir_path / corrupt_name
        corrupt_path.write_bytes(b"%PDF-corrupted invalid byte sequence")
        self.temp_files.append(corrupt_path)

        resp = self.client.post("/documents/extract", json={"saved_filename": corrupt_name})
        self.assertEqual(resp.status_code, 422)


if __name__ == "__main__":
    unittest.main(verbosity=2)
